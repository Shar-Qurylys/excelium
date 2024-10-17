from datetime import datetime
import logging
# from shutil import copy2
from utils.scripts import format_row, set_border, find_last_row_in_col, load_excel, hide_sheets, create_concatenated_info, set_print_area, add_colontituls, set_cell_properties
import utils.firmen_und_objekte as firmobj
from openpyxl.styles import Alignment, Font, Border, Side

def add_coordinators_v4(sheet):
    '''
    Adds coordinators to the specified sheet.

    Args:
        sheet (openpyxl.worksheet.worksheet.Worksheet): The worksheet to add coordinators to.

    Returns:
        None
    '''

    col_index = 6 #F
    formula_f = '=IFERROR(IF(ISNUMBER(VALUE(INDIRECT("B" & ROW()))), VLOOKUP(VALUE(INDIRECT("B" & ROW())), СПР_ПОДПИСАНТОВ!$B$14:$K$100, 9, 0) & " " & VLOOKUP(VALUE(INDIRECT("B" & ROW())), СПР_ПОДПИСАНТОВ!$B$14:$K$100, 7, 0), ""),"")'
    formula_i = '=IFERROR(IF(ISNUMBER(VALUE(INDIRECT("B" & ROW()))), VLOOKUP(VALUE(INDIRECT("B" & ROW())), СПР_ПОДПИСАНТОВ!$B$14:$K$100, 5, 0), ""),"")'

    final_row = find_last_row_in_col(sheet, col_index)

    if final_row:
        logging.info(f"The last non-empty cell in column {chr(64 + col_index)} of sheet '{sheet.title}' is in row {final_row}.")
    else:
        logging.info(f"No non-empty cells found in column {chr(64 + col_index)} of sheet '{sheet.title}'.")
        final_row = 0 # never happens

    company = sheet['H11'].value # Get the company name from the sheet
    object_name = sheet['G11'].value # Get the object name from the sheet
    doctype = sheet['H12'].value

    directors, coordinators_list = firmobj.check_company_object_pair(company, object_name) # Get the coordinators for the company and object
    if sheet['G17'].value.lower() == 'коммерческие расходы':
        coordinators_list = [i for i in coordinators_list if i not in [4, 6]] # remove selected approvers in payments of commercial expenses
    else:
        pass # continue without changing

    # Turned off
    # if sheet['H12'].value.lower() == 'заявка на налоги': # for tax related payments
    #     coordinators_list =firmobj.return_administration_approvers() # approvers always correspond to List 1 approvers

    n = len(coordinators_list) # Get the number of coordinators
    logging.info(f'Company: {company}, Object: {object_name} Number of coordinators: {n}; coordinators: {coordinators_list}')
    for i in range(1,n+1):
        formula_b = f'=INDEX(СПР_ОБЪЕКТОВ!$B$7:$K$80, MATCH($G11, СПР_ОБЪЕКТОВ!$B$7:$B$80, 0), {3 + i})'
        row = final_row + i * 3

        if coordinators_list[i-1] != 3:
            set_cell_properties(sheet, row, 2, coordinators_list[i-1], set_border('thin'))
            set_cell_properties(sheet, row, 6, formula_f, None, Alignment(horizontal='left'), Font(size=14, bold=True))
            set_cell_properties(sheet, row, 9, formula_i, None, Alignment(horizontal='right'), Font(size=14, bold=True))
        else:
            set_cell_properties(sheet, row, 6, "СОГЛАСОВАНО", None, Alignment(horizontal='left'), Font(bold=False))
            set_cell_properties(sheet, row + 2, 2, 3, set_border('thin'))
            set_cell_properties(sheet, row + 2, 6, formula_f, None, Alignment(horizontal='left'), Font(size=14, bold=True))
            set_cell_properties(sheet, row + 2, 9, formula_i, None, Alignment(horizontal='right'), Font(size=14, bold=True))

    # Add directors (final piece)
    sheet['B2'] = directors[0]
    sheet['B4'] = directors[1]
    sheet['H11'] = '' # remove the temporary company name
    sheet['H12'] = '' # remove the temporary document type

def loop_json(json_data, workbook):
    '''
    This function works with the loaded json and with the copied workbook
    '''
    cols = ['F', 'G', 'H', 'I']
    for key_title, data in json_data.items():

        print(len(data)) #how many documents are fetched
        sub = 0 # sub-nomer reestra
        for i in range(len(data)):
            if data[i]['object_name'] not in workbook.sheetnames:
                sub += 1

                start_row = 17 # starting row for the writing
                row = start_row
                source_sheet = workbook['REESTR']

                # Create a copy of the source sheet with the desired name
                new_sheet = workbook.copy_worksheet(source_sheet)

                #listname feststelln
                new_sheet.title = data[i]['object_name']

                #objektname speichern
                object_name = data[i]['object_name']
                workbook[object_name][f'G11'] = object_name
                workbook[object_name][f'G10'] = datetime.today()
                registry_name = data[i]['registry_name']
                workbook[object_name][f'F7'] = f'{registry_name}/{sub}'

                # On dr JSON datei bekommn

                workbook[object_name][f'H11'] = data[i]["organization"]
                workbook[object_name][f'H12'] = data[i]["doctype"]

                sides_str = f'Заявитель: {data[i]["organization"]}'+'\n\n'+f'Кому: {data[i]["counteragent"]}'

                workbook[object_name][f'F{row}'] = sides_str

                workbook[object_name][f'G{row}'] = data[i]['zatraty'] # Zatraty po DDS

                workbook[object_name][f'H{row}'] = float(data[i]['payment_sum']) # Gebuhr

                i_cell_str = create_concatenated_info(data[i])

                workbook[object_name][f'I{row}'] = i_cell_str

                #Zeile formatieren
                format_row(workbook[object_name], row, cols)
                workbook[object_name][f'F{row}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

                start_row += 1
            else:
                # test
                row = start_row
                object_name = data[i]['object_name']

                sides_str = f'Заявитель: {data[i]["organization"]}'+'\n\n'+f'Кому: {data[i]["counteragent"]}'

                workbook[object_name][f'F{row}'] = sides_str

                workbook[object_name][f'G{row}'] = data[i]['zatraty'] # Zatraty po DDS

                workbook[object_name][f'H{row}'] = float(data[i]['payment_sum']) # Gebuhr

                #workbook[object_name][f'M{row}'] = data[i]['sluzhebnaja_zapiska'] # Objektname

                i_cell_str = create_concatenated_info(data[i])
                workbook[object_name][f'I{row}'] = i_cell_str

                #Zeile formatieren
                format_row(workbook[object_name], row, cols)
                workbook[object_name][f'F{row}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

                start_row += 1

def format_excel_inner(json_data):
    logging.info('Opening template.xlsx')
    workbook = load_excel('excel_templates/template.xlsx')
    initial_sheets = ['REESTR', 'СПР_ОБЪЕКТОВ', 'СПР_ПОДПИСАНТОВ']

    logging.info('Removing REESTR sheet')
    if 'REESTR' in workbook.sheetnames:
        reestr_sheet = workbook['REESTR']
        workbook.remove(reestr_sheet)

    logging.info('Reading JSON file')
    loop_json(json_data, workbook)

    # Update the initial_sheets list after removing REESTR
    initial_sheets.remove('REESTR')
    
    hide_sheets(workbook, initial_sheets)
    for sheet in workbook.sheetnames:
        if sheet not in initial_sheets:
            add_coordinators_v4(workbook[sheet])
            set_print_area(workbook[sheet])
            add_colontituls(workbook[sheet])
        else:
            continue

    return workbook
